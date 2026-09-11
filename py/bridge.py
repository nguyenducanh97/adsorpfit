"""
AdsorpFit - JavaScript <-> Python bridge.

Everything the browser calls goes through this module.  All functions take
and return JSON strings, which keeps the Pyodide boundary simple: no proxy
objects leak into JavaScript and no JavaScript objects leak into Python.
"""

from __future__ import annotations

import base64
import io
import json
import math
import traceback

import numpy as np

from core import (ModelSpec, fit_model, fit_linear, rank_models, statistics,
                  R_GAS, fmt)
from isotherms import ISOTHERM_MODELS
from kinetics import KINETIC_MODELS, boyd_bt
import thermo as th
import advisor as adv

ALL_MODELS: dict[str, ModelSpec] = {}
ALL_MODELS.update(ISOTHERM_MODELS)
ALL_MODELS.update(KINETIC_MODELS)


# --------------------------------------------------------------------------
# JSON helpers
# --------------------------------------------------------------------------

def _clean(o):
    """Make any numpy/float value JSON-safe (NaN and Inf are not valid JSON)."""
    if isinstance(o, dict):
        return {k: _clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_clean(v) for v in o]
    if isinstance(o, np.ndarray):
        return [_clean(v) for v in o.tolist()]
    if isinstance(o, (np.floating, float)):
        f = float(o)
        return None if not math.isfinite(f) else f
    if isinstance(o, (np.integer, int)):
        return int(o)
    if isinstance(o, (np.bool_, bool)):
        return bool(o)
    return o


def _ok(payload):
    return json.dumps({"ok": True, **_clean(payload)})


def _err(msg, detail=""):
    return json.dumps({"ok": False, "error": str(msg), "detail": detail})


# --------------------------------------------------------------------------
# model catalogue
# --------------------------------------------------------------------------

def list_models(category: str = "all") -> str:
    """Return the full metadata catalogue for the model picker."""
    out = []
    for key, m in ALL_MODELS.items():
        if category not in ("all", m.category):
            continue
        out.append({
            "key": key, "name": m.name, "category": m.category,
            "family": m.family, "n_params": m.n_params,
            "equation": m.equation, "equation_plain": m.equation_plain,
            "citation": m.citation, "year": m.year,
            "assumptions": m.assumptions,
            "requires": m.requires,
            "notes": m.notes,
            "params": [{
                "key": p.key, "symbol": p.symbol, "unit": p.unit,
                "meaning": p.meaning, "lower": _clean(p.lower),
                "upper": _clean(p.upper),
            } for p in m.params],
            "linear_forms": [{
                "name": lf.name, "x_label": lf.x_label,
                "y_label": lf.y_label, "note": lf.note,
            } for lf in m.linear_forms],
        })
    order = {"2-parameter": 0, "3-parameter": 1, "4-parameter": 2,
             "reaction model": 0, "diffusion model": 1, "empirical model": 2}
    out.sort(key=lambda d: (order.get(d["family"], 9), d["n_params"], d["name"]))
    return _ok({"models": out})


def k_routes() -> str:
    """Return the available routes to a dimensionless equilibrium constant."""
    return _ok({"routes": [{"key": k, **v} for k, v in th.K_ROUTES.items()]})


# --------------------------------------------------------------------------
# fitting
# --------------------------------------------------------------------------

def fit(payload_json: str) -> str:
    """Fit a set of models to one dataset.

    payload = {
        category: "isotherm" | "kinetics",
        x: [...], y: [...],
        models: ["langmuir", ...],
        ctx: {T: 298.15, C0: ..., C0_list: [...], Cs: ..., MW: ...},
        weights: "none" | "relative" | "sqrt",
        also_linear: true/false,
        criterion: "AICc"
    }
    """
    try:
        p = json.loads(payload_json)
        x = np.asarray(p["x"], float)
        y = np.asarray(p["y"], float)
        ctx = p.get("ctx", {}) or {}
        weights = p.get("weights", "none")
        also_linear = bool(p.get("also_linear", True))
        criterion = p.get("criterion", "AICc")

        if x.size != y.size:
            return _err(f"x has {x.size} values but y has {y.size}. "
                        f"Every point needs both a concentration/time and a q value.")
        if x.size < 3:
            return _err("At least 3 data points are needed to fit anything "
                        "meaningfully.")

        results, payload_models = [], []
        for key in p["models"]:
            spec = ALL_MODELS.get(key)
            if spec is None:
                continue
            missing = [r for r in spec.requires if r not in ctx]
            r = fit_model(spec, x, y, ctx=ctx, weights=weights)
            results.append(r)

            entry = _result_to_dict(r, spec, ctx)
            if missing:
                entry["warnings"] = entry.get("warnings", []) + [
                    f"This model needs {', '.join(missing)} from the experiment "
                    f"panel. A default was used, so the parameters are not on a "
                    f"physical scale."]

            if also_linear and spec.linear_forms:
                entry["linear"] = []
                for lf in spec.linear_forms:
                    lr = fit_linear(spec, lf, x, y, ctx=ctx)
                    d = _result_to_dict(lr, spec, ctx, with_curve=False)
                    d["form_name"] = lf.name
                    d["x_label"] = lf.x_label
                    d["y_label"] = lf.y_label
                    d["note"] = lf.note
                    if lr.success:
                        xs = lr.derived.get("linear_x")
                        ys = lr.derived.get("linear_y")
                        msk = lr.derived.get("linear_mask")
                        if xs is not None:
                            d["plot"] = {
                                "x": _clean(np.asarray(xs)[msk]),
                                "y": _clean(np.asarray(ys)[msk]),
                                "slope": _clean(lr.stats.get("slope")),
                                "intercept": _clean(lr.stats.get("intercept")),
                                "R2": _clean(lr.stats.get("R2_linear")),
                            }
                    entry["linear"].append(d)
            payload_models.append(entry)

        ranking = rank_models(results, criterion=criterion)
        return _ok({
            "results": payload_models,
            "ranking": _clean(ranking),
            "summary": _ranking_prose(ranking, results),
        })
    except Exception as exc:
        return _err(exc, traceback.format_exc())


def advise(payload_json: str) -> str:
    """Screen every model against the physics of the raw data.

    payload = {category, x, y, ctx, models?}
    """
    try:
        p = json.loads(payload_json)
        out = adv.advise(p["category"], p["x"], p["y"],
                         ctx=p.get("ctx", {}) or {},
                         models=p.get("models"))
        return _ok(out)
    except Exception as exc:
        return _err(exc, traceback.format_exc())


def _result_to_dict(r, spec, ctx, with_curve=True):
    d = {
        "model_key": r.model_key, "model_name": r.model_name,
        "method": r.method, "success": r.success, "message": r.message,
        "params": _clean(r.params), "stderr": _clean(r.stderr),
        "ci95": {k: _clean(list(v)) for k, v in r.ci95.items()},
        "tvalue": _clean(r.tvalue), "pvalue": _clean(r.pvalue),
        "stats": _clean(r.stats),
        "residuals": _clean(r.residuals),
        "y_cal": _clean(r.y_cal),
        "warnings": list(r.warnings),
        "issues": list(r.issues),
        "equation": spec.equation, "equation_plain": spec.equation_plain,
        "citation": spec.citation, "assumptions": spec.assumptions,
        "family": spec.family, "category": spec.category,
        "param_meta": [{"key": p.key, "symbol": p.symbol, "unit": p.unit,
                        "meaning": p.meaning} for p in spec.params],
    }
    if r.success and with_curve and r._fn is not None:
        try:
            lo = 0.0
            hi = float(np.max(r.x)) * 1.05
            xs = np.linspace(lo, hi, 300)
            ys = r._fn(xs)
            good = np.isfinite(ys)
            d["curve"] = {"x": _clean(xs[good]), "y": _clean(ys[good])}
        except Exception:
            d["curve"] = None
    if r.success and spec.interpretation is not None:
        try:
            d["interpretation"] = [s for s in spec.interpretation(r, ctx) if s]
        except Exception as exc:
            d["interpretation"] = [f"(interpretation unavailable: {exc})"]
    d["derived"] = _clean({k: v for k, v in r.derived.items()
                           if not isinstance(v, np.ndarray) or v.ndim <= 1})
    return d


def _ranking_prose(ranking, results):
    """Write the model-comparison paragraph."""
    if not ranking:
        return ["No model converged, so there is nothing to rank."]
    best = ranking[0]
    lines = [
        f"**{best['model_name']}** ranks first on {best['criterion']} "
        f"({best['criterion']} = {best['value']:.2f}, Akaike weight "
        f"{best['weight'] * 100:.0f}%), with R² = {best['R2']:.4f} and RMSE = "
        f"{best['RMSE']:.4g}."
    ]
    close = [r for r in ranking[1:] if r["delta"] < 2]
    if close:
        names = ", ".join(r["model_name"] for r in close)
        lines.append(
            f"However, {names} {'is' if len(close) == 1 else 'are'} statistically "
            f"indistinguishable from it (ΔAICc < 2). On these data you cannot "
            f"claim one of these models is correct and the others are not, say "
            f"they describe the data equally well, and choose between them on "
            f"physical grounds rather than on fit statistics."
        )
    else:
        second = ranking[1] if len(ranking) > 1 else None
        if second:
            lines.append(
                f"The next best model, {second['model_name']}, is Δ = "
                f"{second['delta']:.1f} behind: {second['evidence']}. The margin "
                f"is large enough to prefer {best['model_name']} on statistical "
                f"grounds."
            )
    # the R2 trap
    by_r2 = sorted(ranking, key=lambda r: -(r["R2"] if r["R2"] is not None else -9))
    if by_r2[0]["model_key"] != best["model_key"]:
        lines.append(
            f"Note that {by_r2[0]['model_name']} has the highest raw R² "
            f"({by_r2[0]['R2']:.4f}) but uses {by_r2[0]['n_params']} parameters "
            f"against {best['n_params']} for {best['model_name']}. R² can only "
            f"increase when parameters are added, so ranking models by R² "
            f"systematically favours the most complex one. AICc penalises that "
            f"and is the criterion to report."
        )
    return lines


# --------------------------------------------------------------------------
# diffusion diagnostics
# --------------------------------------------------------------------------

def diffusion_analysis(payload_json: str) -> str:
    """Weber-Morris multi-region segmentation and the Boyd plot."""
    try:
        p = json.loads(payload_json)
        t = np.asarray(p["t"], float)
        q = np.asarray(p["q"], float)
        n_seg = int(p.get("segments", 2))
        qe = p.get("qe") or float(np.max(q))

        order = np.argsort(t)
        t, q = t[order], q[order]
        sq = np.sqrt(np.clip(t, 0, None))

        segments = _segment_weber_morris(sq, q, n_seg)
        F = np.clip(q / qe, 1e-9, 1 - 1e-9)
        bt = boyd_bt(F)
        ok = np.isfinite(bt) & (t > 0)
        from scipy import stats as sps
        boyd = {}
        if ok.sum() >= 2:
            lr = sps.linregress(t[ok], bt[ok])
            boyd = {
                "slope": lr.slope, "intercept": lr.intercept,
                "R2": lr.rvalue ** 2,
                "t": _clean(t[ok]), "Bt": _clean(bt[ok]),
                "se_intercept": _clean(lr.intercept_stderr),
            }
            # is the intercept significantly different from zero?
            sig = (abs(lr.intercept) > 2 * lr.intercept_stderr
                   if np.isfinite(lr.intercept_stderr) else abs(lr.intercept) > 0.1)
            boyd["through_origin"] = not sig
            boyd["interpretation"] = (
                f"The Boyd plot has intercept {fmt(lr.intercept)} ± "
                f"{fmt(lr.intercept_stderr)} and R² = {lr.rvalue ** 2:.4f}. "
                + ("The intercept is not statistically distinguishable from zero, "
                   "so the line passes through the origin: **intraparticle "
                   "(pore) diffusion** controls the rate."
                   if not sig else
                   "The intercept differs significantly from zero, so the line "
                   "does NOT pass through the origin: **film diffusion** (transport "
                   "across the liquid boundary layer around the particle), or a "
                   "chemical reaction step, controls the rate rather than "
                   "intraparticle diffusion.")
            )
            if np.isfinite(lr.slope) and lr.slope > 0:
                r = p.get("particle_radius")
                if r:
                    Di = lr.slope * (float(r) ** 2) / (np.pi ** 2)
                    boyd["Di"] = Di
                    boyd["interpretation"] += (
                        f" From the slope B = {fmt(lr.slope)} min⁻¹ and a particle "
                        f"radius of {r} cm, the effective diffusion coefficient is "
                        f"D_i = B·r²/π² = {fmt(Di)} cm²/min.")
        return _ok({"weber_morris": segments, "boyd": _clean(boyd)})
    except Exception as exc:
        return _err(exc, traceback.format_exc())


def _segment_weber_morris(sq, q, n_seg):
    """Split the qt vs sqrt(t) plot into n_seg linear regions.

    Breakpoints are chosen by exhaustive search over all admissible splits,
    minimising total residual sum of squares.  With the small point counts of
    a kinetics experiment this is cheap and avoids the arbitrariness of
    choosing breakpoints by eye - which is how it is normally done.
    """
    from scipy import stats as sps
    n = sq.size
    n_seg = max(1, min(n_seg, max(1, n // 3)))

    def seg_fit(i0, i1):
        xs, ys = sq[i0:i1], q[i0:i1]
        if xs.size < 2:
            return None, np.inf
        lr = sps.linregress(xs, ys)
        pred = lr.slope * xs + lr.intercept
        return lr, float(np.sum((ys - pred) ** 2))

    best, best_sse = None, np.inf
    if n_seg == 1:
        lr, sse = seg_fit(0, n)
        best, best_sse = [(0, n, lr)], sse
    elif n_seg == 2:
        for b in range(2, n - 1):
            l1, s1 = seg_fit(0, b)
            l2, s2 = seg_fit(b, n)
            if l1 is None or l2 is None:
                continue
            if s1 + s2 < best_sse:
                best, best_sse = [(0, b, l1), (b, n, l2)], s1 + s2
    else:
        for b1 in range(2, n - 3):
            for b2 in range(b1 + 2, n - 1):
                l1, s1 = seg_fit(0, b1)
                l2, s2 = seg_fit(b1, b2)
                l3, s3 = seg_fit(b2, n)
                if None in (l1, l2, l3):
                    continue
                if s1 + s2 + s3 < best_sse:
                    best = [(0, b1, l1), (b1, b2, l2), (b2, n, l3)]
                    best_sse = s1 + s2 + s3

    if best is None:
        lr, sse = seg_fit(0, n)
        best = [(0, n, lr)]

    labels = {
        1: ["external surface adsorption (instantaneous stage)"],
        2: ["stage 1: external surface adsorption / film diffusion",
            "stage 2: intraparticle diffusion (rate-limiting)"],
        3: ["stage 1: external surface adsorption / film diffusion",
            "stage 2: intraparticle diffusion (rate-limiting)",
            "stage 3: equilibrium plateau, diffusion slows as sites fill"],
    }.get(len(best), [f"stage {i + 1}" for i in range(len(best))])

    out = []
    for i, (i0, i1, lr) in enumerate(best):
        out.append({
            "stage": i + 1,
            "label": labels[i] if i < len(labels) else f"stage {i + 1}",
            "kid": _clean(lr.slope), "C": _clean(lr.intercept),
            "R2": _clean(lr.rvalue ** 2),
            "kid_se": _clean(lr.stderr), "C_se": _clean(lr.intercept_stderr),
            "x": _clean(sq[i0:i1]), "y": _clean(q[i0:i1]),
            "n_points": int(i1 - i0),
        })
    if len(out) >= 2:
        slower = out[0]["kid"] > out[1]["kid"]
        out_note = (
            f"The plot resolves into {len(out)} linear regions, which is the "
            f"expected multi-step behaviour. Region 1 has k_id = "
            f"{fmt(out[0]['kid'])} and region 2 has k_id = {fmt(out[1]['kid'])} "
            f"mg g⁻¹ min⁻⁰·⁵. "
            + ("The first stage is faster, as expected: adsorbate first covers the "
               "readily accessible external surface, then diffusion into the pores "
               "takes over and slows the process."
               if slower else
               "Unusually, the second stage is faster than the first. Check the "
               "segmentation: this can happen when the first region contains too "
               "few points to define a slope."))
    else:
        out_note = ("Only one linear region was resolved. If the whole plot really "
                    "is a single straight line through the origin, intraparticle "
                    "diffusion controls the rate throughout.")
    return {"segments": out, "note": out_note}


# --------------------------------------------------------------------------
# thermodynamics
# --------------------------------------------------------------------------

def fit_thermo(payload_json: str) -> str:
    """Full thermodynamic analysis.

    payload = {
        T: [K], route: "langmuir_molar", values: {...per-temperature...},
        nonlinear: bool, MW: float,
        arrhenius: {T: [...], k: [...]},
        isosteric: {T: [...], q: [...], Ce: [[...]]}
    }
    """
    try:
        p = json.loads(payload_json)
        T = np.asarray(p["T"], float)
        route = p.get("route", "langmuir_molar")
        raw = p.get("values", [])

        K0, details, warns = [], [], []
        for i, v in enumerate(raw):
            vals = dict(v)
            if route == "langmuir_molar" and "MW" not in vals:
                vals["MW"] = p.get("MW", 1.0)
            try:
                conv = th.dimensionless_K(route, vals)
            except Exception as exc:
                return _err(f"K° conversion failed at T = {T[i]} K: {exc}")
            K0.append(conv["K0"])
            details.append(conv["detail"])
            warns.extend(conv["warnings"])

        res = th.vant_hoff(T, np.asarray(K0, float),
                           nonlinear=bool(p.get("nonlinear", False)))
        route_info = {"route": route, "detail": "; ".join(details[:1]) +
                      (f" (and {len(details) - 1} more temperatures)"
                       if len(details) > 1 else "")}
        out = {"vant_hoff": _clean(res),
               "K0": _clean(K0),
               "route": route,
               "route_label": th.K_ROUTES[route]["label"],
               "route_note": th.K_ROUTES[route]["note"],
               "route_defensible": th.K_ROUTES[route]["defensible"],
               "conversion_details": details,
               "interpretation": th.interpret_thermo(res, route_info),
               "warnings": list(dict.fromkeys(warns))}

        if p.get("arrhenius"):
            a = p["arrhenius"]
            out["arrhenius"] = _clean(th.arrhenius(a["T"], a["k"]))
        if p.get("isosteric"):
            iso = p["isosteric"]
            out["isosteric"] = _clean(
                th.isosteric_heat(iso["T"], iso["Ce"], iso["q"]))
        if p.get("sticking"):
            s = p["sticking"]
            out["sticking"] = _clean(th.sticking_probability(s["T"], s["theta"]))
        return _ok(out)
    except Exception as exc:
        return _err(exc, traceback.format_exc())


def isotherm_K_for_thermo(payload_json: str) -> str:
    """Fit one isotherm model at each temperature and return its constants.

    This is the step most papers do by hand: fit Langmuir (or R-P) at every
    temperature, then carry K_L into the van't Hoff plot.
    """
    try:
        p = json.loads(payload_json)
        model_key = p.get("model", "langmuir")
        spec = ALL_MODELS[model_key]
        rows = []
        for ds in p["datasets"]:
            T = float(ds["T"])
            ctx = dict(p.get("ctx", {}))
            ctx["T"] = T
            r = fit_model(spec, np.asarray(ds["x"], float),
                          np.asarray(ds["y"], float), ctx=ctx,
                          weights=p.get("weights", "none"))
            rows.append({
                "T": T, "success": r.success,
                "params": _clean(r.params), "stderr": _clean(r.stderr),
                "R2": _clean(r.stats.get("R2")),
                "AICc": _clean(r.stats.get("AICc")),
            })
        return _ok({"model": model_key, "model_name": spec.name, "rows": rows})
    except Exception as exc:
        return _err(exc, traceback.format_exc())


# --------------------------------------------------------------------------
# publication-quality figure rendering (matplotlib)
# --------------------------------------------------------------------------

def render_figure(payload_json: str) -> str:
    """Render a figure with matplotlib and return it base64-encoded.

    Supported formats: png, pdf, svg, eps, ps, jpg, tiff, webp.
    matplotlib writes the first five natively; tiff, jpg and webp go through
    Pillow from a high-resolution PNG buffer so the requested DPI is honoured.
    """
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.ticker import AutoMinorLocator

        p = json.loads(payload_json)
        s = p.get("style", {})
        fmt_ = p.get("format", "png").lower()
        dpi = int(p.get("dpi", 600))

        width = float(s.get("width_cm", 8.5)) / 2.54
        height = float(s.get("height_cm", 6.5)) / 2.54
        fig, ax = plt.subplots(figsize=(width, height), dpi=dpi)

        font = s.get("font_family", "DejaVu Sans")
        fs = float(s.get("font_size", 9))
        plt.rcParams["font.family"] = font
        plt.rcParams["mathtext.fontset"] = "dejavusans"

        for tr in p.get("traces", []):
            kind = tr.get("kind", "scatter")
            if kind == "scatter":
                ax.errorbar(
                    tr["x"], tr["y"],
                    yerr=tr.get("yerr"), xerr=tr.get("xerr"),
                    fmt=tr.get("marker", "o"),
                    markersize=float(tr.get("marker_size", 5)),
                    markerfacecolor=tr.get("marker_fill", tr.get("color", "#0a6ebd")),
                    markeredgecolor=tr.get("marker_edge", tr.get("color", "#0a6ebd")),
                    markeredgewidth=float(tr.get("marker_edge_width", 1.0)),
                    color=tr.get("color", "#0a6ebd"),
                    ecolor=tr.get("error_color", tr.get("color", "#0a6ebd")),
                    elinewidth=float(tr.get("error_width", 0.8)),
                    capsize=float(tr.get("capsize", 2.5)),
                    linestyle="none", label=tr.get("name"),
                    zorder=tr.get("zorder", 3),
                )
            else:
                ax.plot(
                    tr["x"], tr["y"],
                    color=tr.get("color", "#c1272d"),
                    linewidth=float(tr.get("line_width", 1.4)),
                    linestyle=tr.get("line_style", "-"),
                    label=tr.get("name"),
                    zorder=tr.get("zorder", 2),
                )

        ax.set_xlabel(s.get("x_label", ""), fontsize=fs,
                      fontweight=s.get("label_weight", "normal"))
        ax.set_ylabel(s.get("y_label", ""), fontsize=fs,
                      fontweight=s.get("label_weight", "normal"))
        if s.get("title"):
            ax.set_title(s["title"], fontsize=fs * 1.1,
                         fontweight=s.get("title_weight", "normal"))

        ax.tick_params(axis="both", which="major",
                       labelsize=fs * float(s.get("tick_scale", 0.92)),
                       direction=s.get("tick_direction", "in"),
                       length=float(s.get("tick_length", 4)),
                       width=float(s.get("axis_width", 0.9)),
                       top=bool(s.get("mirror_ticks", True)),
                       right=bool(s.get("mirror_ticks", True)))
        if s.get("minor_ticks", True):
            ax.xaxis.set_minor_locator(AutoMinorLocator())
            ax.yaxis.set_minor_locator(AutoMinorLocator())
            ax.tick_params(axis="both", which="minor",
                           direction=s.get("tick_direction", "in"),
                           length=float(s.get("tick_length", 4)) * 0.55,
                           width=float(s.get("axis_width", 0.9)) * 0.8,
                           top=bool(s.get("mirror_ticks", True)),
                           right=bool(s.get("mirror_ticks", True)))

        for side in ("top", "right", "bottom", "left"):
            ax.spines[side].set_linewidth(float(s.get("axis_width", 0.9)))
            if side in ("top", "right") and not s.get("box", True):
                ax.spines[side].set_visible(False)

        if s.get("grid"):
            ax.grid(True, which=s.get("grid_which", "major"),
                    linestyle=s.get("grid_style", ":"),
                    linewidth=float(s.get("grid_width", 0.6)),
                    color=s.get("grid_color", "#b8c4cc"),
                    alpha=float(s.get("grid_alpha", 0.8)), zorder=0)

        if s.get("x_log"):
            ax.set_xscale("log")
        if s.get("y_log"):
            ax.set_yscale("log")
        if s.get("x_min") is not None and s.get("x_max") is not None:
            ax.set_xlim(float(s["x_min"]), float(s["x_max"]))
        if s.get("y_min") is not None and s.get("y_max") is not None:
            ax.set_ylim(float(s["y_min"]), float(s["y_max"]))

        if s.get("legend", True):
            leg = ax.legend(
                loc=s.get("legend_loc", "lower right"),
                fontsize=fs * float(s.get("legend_scale", 0.88)),
                frameon=bool(s.get("legend_frame", False)),
                ncol=int(s.get("legend_cols", 1)),
                handlelength=float(s.get("legend_handle", 1.6)),
                labelspacing=float(s.get("legend_spacing", 0.35)),
                borderpad=0.4,
            )
            if leg and s.get("legend_frame"):
                leg.get_frame().set_linewidth(0.6)
                leg.get_frame().set_edgecolor(s.get("legend_edge", "#333333"))

        if s.get("annotation"):
            ax.annotate(s["annotation"],
                        xy=(float(s.get("ann_x", 0.05)), float(s.get("ann_y", 0.92))),
                        xycoords="axes fraction", fontsize=fs * 0.9,
                        va="top", ha="left")

        fig.tight_layout(pad=float(s.get("pad", 0.3)))

        buf = io.BytesIO()
        if fmt_ in ("png", "pdf", "svg", "eps", "ps"):
            fig.savefig(buf, format=fmt_, dpi=dpi,
                        bbox_inches="tight" if s.get("tight", True) else None,
                        transparent=bool(s.get("transparent", False)),
                        facecolor="none" if s.get("transparent") else "white")
            mime = {"png": "image/png", "pdf": "application/pdf",
                    "svg": "image/svg+xml", "eps": "application/postscript",
                    "ps": "application/postscript"}[fmt_]
        else:
            png = io.BytesIO()
            fig.savefig(png, format="png", dpi=dpi,
                        bbox_inches="tight" if s.get("tight", True) else None,
                        facecolor="white")
            png.seek(0)
            from PIL import Image
            im = Image.open(png)
            if fmt_ in ("jpg", "jpeg"):
                im = im.convert("RGB")
                im.save(buf, format="JPEG", quality=int(s.get("quality", 95)),
                        dpi=(dpi, dpi))
                mime = "image/jpeg"
            elif fmt_ in ("tif", "tiff"):
                im = im.convert("RGB")
                im.save(buf, format="TIFF",
                        compression=s.get("tiff_compression", "tiff_lzw"),
                        dpi=(dpi, dpi))
                mime = "image/tiff"
            elif fmt_ == "webp":
                im.save(buf, format="WEBP", quality=int(s.get("quality", 95)))
                mime = "image/webp"
            elif fmt_ == "bmp":
                im = im.convert("RGB")
                im.save(buf, format="BMP")
                mime = "image/bmp"
            else:
                plt.close(fig)
                return _err(f"Unsupported export format: {fmt_}")

        plt.close(fig)
        buf.seek(0)
        data = base64.b64encode(buf.read()).decode("ascii")
        return _ok({"data": data, "mime": mime, "format": fmt_, "dpi": dpi})
    except Exception as exc:
        return _err(exc, traceback.format_exc())


def available_fonts() -> str:
    """List the font families matplotlib can actually use in this runtime."""
    try:
        import matplotlib.font_manager as fm
        names = sorted({f.name for f in fm.fontManager.ttflist})
        return _ok({"fonts": names})
    except Exception as exc:
        return _err(exc)


# --------------------------------------------------------------------------
# data export
# --------------------------------------------------------------------------

def export_table(payload_json: str) -> str:
    """Produce a results table in csv / tsv / latex / markdown / html format."""
    try:
        p = json.loads(payload_json)
        rows = p["rows"]
        cols = p["columns"]
        kind = p.get("format", "csv")
        cap = p.get("caption", "Fitted model parameters")

        def cell(r, c):
            v = r.get(c["key"], "")
            if isinstance(v, float):
                return fmt(v, int(p.get("sig", 4)))
            return "" if v is None else str(v)

        if kind in ("csv", "tsv"):
            sep = "," if kind == "csv" else "\t"
            out = [sep.join(c["label"] for c in cols)]
            out += [sep.join(cell(r, c) for c in cols) for r in rows]
            text = "\n".join(out)
        elif kind == "markdown":
            out = ["| " + " | ".join(c["label"] for c in cols) + " |",
                   "|" + "|".join("---" for _ in cols) + "|"]
            out += ["| " + " | ".join(cell(r, c) for c in cols) + " |" for r in rows]
            text = "\n".join(out)
        elif kind == "latex":
            align = "l" + "r" * (len(cols) - 1)
            out = [r"\begin{table}[htbp]", r"\centering",
                   rf"\caption{{{cap}}}",
                   rf"\begin{{tabular}}{{{align}}}", r"\hline",
                   " & ".join(c["label"] for c in cols) + r" \\", r"\hline"]
            out += [" & ".join(cell(r, c) for c in cols) + r" \\" for r in rows]
            out += [r"\hline", r"\end{tabular}", r"\end{table}"]
            text = "\n".join(out)
        elif kind == "html":
            out = ["<table>", "<caption>" + cap + "</caption>", "<thead><tr>"]
            out += [f"<th>{c['label']}</th>" for c in cols]
            out += ["</tr></thead><tbody>"]
            for r in rows:
                out.append("<tr>" + "".join(f"<td>{cell(r, c)}</td>"
                                            for c in cols) + "</tr>")
            out += ["</tbody></table>"]
            text = "\n".join(out)
        else:
            return _err(f"Unknown table format: {kind}")
        return _ok({"text": text, "format": kind})
    except Exception as exc:
        return _err(exc, traceback.format_exc())


def export_xlsx(payload_json: str) -> str:
    """Write a multi-sheet .xlsx workbook and return it base64-encoded."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill

        p = json.loads(payload_json)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        head_fill = PatternFill("solid", fgColor="1B6CA8")
        head_font = Font(color="FFFFFF", bold=True)

        for sheet in p["sheets"]:
            ws = wb.create_sheet(title=sheet["name"][:31])
            for j, h in enumerate(sheet["columns"], start=1):
                c = ws.cell(row=1, column=j, value=h)
                c.fill = head_fill
                c.font = head_font
                c.alignment = Alignment(horizontal="center")
            for i, row in enumerate(sheet["rows"], start=2):
                for j, v in enumerate(row, start=1):
                    ws.cell(row=i, column=j,
                            value=(v if not isinstance(v, float)
                                   or math.isfinite(v) else None))
            for j, h in enumerate(sheet["columns"], start=1):
                width = max(10, min(40, len(str(h)) + 4))
                ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = width
            ws.freeze_panes = "A2"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return _ok({"data": base64.b64encode(buf.read()).decode("ascii"),
                    "mime": "application/vnd.openxmlformats-officedocument."
                            "spreadsheetml.sheet"})
    except Exception as exc:
        return _err(exc, traceback.format_exc())


def parse_pasted(text: str) -> str:
    """Parse pasted spreadsheet data into numeric columns.

    Accepts tab, comma, semicolon or whitespace separation, tolerates a
    header row, European decimal commas, and thousands separators.
    """
    try:
        lines = [ln for ln in text.replace("\r", "").split("\n") if ln.strip()]
        if not lines:
            return _err("Nothing to parse.")

        def split(ln):
            for sep in ("\t", ";", ","):
                if sep in ln:
                    return [c.strip() for c in ln.split(sep)]
            return ln.split()

        rows = [split(ln) for ln in lines]
        header = None
        try:
            [float(c.replace(",", ".")) for c in rows[0] if c]
        except ValueError:
            header = rows[0]
            rows = rows[1:]

        ncol = max(len(r) for r in rows)
        cols = [[] for _ in range(ncol)]
        skipped = 0
        for r in rows:
            if len(r) < ncol:
                r = r + [""] * (ncol - len(r))
            vals = []
            bad = False
            for c in r[:ncol]:
                c = c.strip().replace("−", "-")
                if not c:
                    vals.append(None)
                    continue
                try:
                    vals.append(float(c))
                except ValueError:
                    try:
                        vals.append(float(c.replace(".", "").replace(",", ".")))
                    except ValueError:
                        bad = True
                        break
            if bad:
                skipped += 1
                continue
            for i, v in enumerate(vals):
                cols[i].append(v)

        return _ok({"columns": _clean(cols), "header": header,
                    "n_rows": len(cols[0]) if cols else 0,
                    "n_cols": ncol, "skipped": skipped})
    except Exception as exc:
        return _err(exc, traceback.format_exc())
